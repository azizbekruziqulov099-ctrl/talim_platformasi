import io
import unittest

import openpyxl

from modules.ai_miya_oddiy import template_workbook as simple_template
from modules.ai_miya_varoq import (COLUMNS, SAMPLE_ROWS, is_sheet_workbook, kod_norm, parse_options, parse_sheets,
                                   split_voice, template_workbook)
from modules.dars_export import lesson_docx, lesson_pdf
from modules.dars_xonasi import build_lesson, practice_item, solution_steps

KEYS = [c[0] for c in COLUMNS]


def filled(sheets, book=None):
    """sheets: [(meta, rows)] → har biri alohida mavzu varag'i."""
    wb = template_workbook(blank_topics=len(sheets), book=book or {"kitob_nomi": "Xalqaro baholash", "fan": "Pedagogika", "sinf": "3 kurs", "kod_prefiksi": "XB"})
    for i, (meta, rows) in enumerate(sheets, 1):
        ws = wb[f"{i}-mavzu"]
        ws.cell(2, 2).value = meta.get("mavzu_kodi")
        ws.cell(2, 7).value = meta.get("mavzu_nomi")
        if meta.get("mavzu_raqami"):
            ws.cell(3, 2).value = meta["mavzu_raqami"]
        for r, row in enumerate(rows, 7):
            for c, key in enumerate(KEYS, 1):
                if key != "tartib" and row.get(key):
                    ws.cell(r, c).value = row[key]
    out = io.BytesIO(); wb.save(out)
    return openpyxl.load_workbook(io.BytesIO(out.getvalue()))


def units_of(p):
    units = [{"unit_code": r["step_id"], "unit_kind": "lesson_step", "payload": r} for r in p["11_DARS_SSENARIY"]]
    units += [{"unit_code": r["variant_id"], "unit_kind": "variant", "payload": r} for r in p["12_TUSHUNMADIM"]]
    units += [{"unit_code": r["task_id"], "unit_kind": "task", "payload": r} for r in p["06_MASHQLAR"]]
    return units


class VaroqliShablonTests(unittest.TestCase):
    def test_template_one_sheet_per_topic_and_sample_is_skipped(self):
        wb = template_workbook([{"mavzu_kodi": "T1", "mavzu_nomi": "Kasr"}, {"mavzu_kodi": "T2", "mavzu_nomi": "Nisbat: [1]"}])
        self.assertEqual(wb.sheetnames, ["KITOB", "NAMUNA", "01 Kasr", "02 Nisbat 1"])
        self.assertTrue(is_sheet_workbook(wb))
        self.assertFalse(is_sheet_workbook(simple_template()))
        parsed = parse_sheets(wb)
        self.assertTrue(any("to'ldirilgan mavzu varag'i topilmadi" in e["message"] for e in parsed["errors"]))

    def test_concepts_then_practice_with_book_codes(self):
        parsed = parse_sheets(filled([({"mavzu_kodi": "5-01", "mavzu_nomi": "Oddiy kasr", "mavzu_raqami": "3"}, SAMPLE_ROWS)]), {"pitsa_8.png"})
        self.assertEqual(parsed["errors"], [])
        p = parsed["payload"]
        kinds = [s["qadam_turi"] for s in p["11_DARS_SSENARIY"]]
        self.assertEqual(kinds, ["qoida", "qoida", "xulosa", "amaliy", "amaliy", "amaliy"])
        codes = [c["kod"] for c in parsed["kodlar"]]
        self.assertEqual(codes, ["XB-03-M01", "XB-03-S01", "XB-03-A01", "XB-03-T01"])
        test = p["06_MASHQLAR"][0]
        self.assertEqual((test["togri_javob"], test["variant_b"], test["kitob_kodi"]), ("B", "9", "XB-03-T01"))
        self.assertEqual(len(p["12_TUSHUNMADIM"]), 2)
        self.assertEqual(parsed["summary"]["kitob_kodlari"], 4)

    def test_duplicate_code_and_bad_test_are_reported(self):
        rows = [{"turi": "tushuncha", "sarlavha": "A", "matn": "Matn."},
                {"turi": "masala", "kod": "A01", "matn": "1-shart", "yechim": "x"},
                {"turi": "topshiriq", "kod": "a01", "matn": "2-shart", "yechim": "y"},
                {"turi": "test", "matn": "Savol?", "variantlar": "A) 1", "javob": "A"}]
        parsed = parse_sheets(filled([({"mavzu_kodi": "K", "mavzu_nomi": "M"}, rows)]))
        self.assertTrue(any("takrorlangan" in e["message"] for e in parsed["errors"]))
        # REV80: 2–4 variant qabul qilinadi (bog'cha testlari 2–3 ta); 1 ta variant — masala bo'lib qoladi.
        self.assertTrue(any("2–4 ta variant" in w["message"] for w in parsed["warnings"]))
        self.assertEqual(parsed["payload"]["06_MASHQLAR"], [])

    def test_topic_without_code_waits_for_name_lookup(self):
        parsed = parse_sheets(filled([({"mavzu_nomi": "PISA dasturi"}, [{"turi": "tushuncha", "sarlavha": "PISA", "matn": "PISA — xalqaro baholash."}])]))
        self.assertEqual(list(parsed["nom_boyicha"].values())[0]["nom"], "PISA dasturi")
        self.assertTrue(parsed["payload"]["02_DTS_XARITA"][0]["topic_code"].startswith("@@NOM"))

    def test_long_explanation_is_split_and_board_lines_follow_markers(self):
        para = "Bu juda muhim gap. " * 40
        text = f"[1] {para} [2] {para} [3] {para}"
        pieces = split_voice(text, ["bir", "ikki", "uch"])
        self.assertGreater(len(pieces), 1)
        self.assertTrue(all(len(v) <= 1500 for v, _ in pieces))
        self.assertEqual("\n".join(b for _, b in pieces if b).split("\n"), ["bir", "ikki", "uch"])
        self.assertTrue(pieces[0][0].startswith("[1]"))

    def test_lesson_board_and_code_item(self):
        parsed = parse_sheets(filled([({"mavzu_kodi": "5-01", "mavzu_nomi": "Oddiy kasr"}, SAMPLE_ROWS)]), {"pitsa_8.png"})
        units = units_of(parsed["payload"])
        lesson = build_lesson({"mavzu": "Oddiy kasr"}, units, lambda m: None)
        practice = [s for s in lesson["steps"] if s["turi"] == "amaliy"]
        self.assertEqual(practice[0]["kod"], "XB-01-M01")
        self.assertEqual(practice[0]["yechim"][0], {"doska": "Maxraj = 8", "ovoz": "Jami sakkiz bo'lak — maxrajga sakkiz yozamiz."})
        self.assertEqual(lesson["savollar"][0]["kod"], "XB-01-T01")
        task = next(u for u in units if u["unit_kind"] == "task")
        item = practice_item(task, lambda m: None)
        self.assertEqual((item["turi"], item["togri"], item["yechim"][0]["doska"]), ("test", 1, "Javob: B) 9"))
        masala = next(u for u in units if u["payload"].get("amaliy_turi") == "masala")
        self.assertEqual(practice_item(masala, lambda m: None)["javob"], "2/6|1/3")
        self.assertIn(b"%PDF", lesson_pdf(lesson, {})[:8])
        self.assertGreater(len(lesson_docx(lesson, {})), 1000)

    def test_helpers(self):
        self.assertEqual(kod_norm(" xb-03 a01 "), "XB03A01")
        self.assertEqual(parse_options("A) 6 km/soat.   B) 18 km/soat.   C) 25"), ["6 km/soat.", "18 km/soat.", "25"])
        self.assertEqual(solution_steps("a || b\nc"), [{"doska": "a", "ovoz": "b"}, {"doska": "c", "ovoz": "c"}])


if __name__ == "__main__":
    unittest.main()

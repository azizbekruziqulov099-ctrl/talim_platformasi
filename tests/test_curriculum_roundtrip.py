"""Generate real XLSX files, fill them, then run the real import handlers.

SQLite executes the relational work; a small execute_values adapter replaces
psycopg2's bulk transport. This is not a PostgreSQL/HTTP integration test.
"""
import ast,asyncio,io,re,sqlite3,sys,unicodedata,unittest
from pathlib import Path
from types import SimpleNamespace,ModuleType
from typing import Optional
from unittest.mock import patch
import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from test_curriculum_scope import DB,HTTPException,scope
from modules.test_template_import import discover_test_worksheets

class Download:
 def __init__(self,body,**kwargs):self.body=body.getvalue();self.headers=kwargs.get('headers',{})

class Upload:
 def __init__(self,workbook):
  self.buffer=io.BytesIO();workbook.save(self.buffer);self.buffer.seek(0)
 async def read(self,size=-1):return self.buffer.read(size)
 async def seek(self,offset):return self.buffer.seek(offset)

class WorkbookRoundTripTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  path=Path(__file__).resolve().parents[1]/'samtm_platform.py'
  names={'tushuntirish_import','mavzu_tushuntirish_ol','topik_royxat','topik_fanlar','topik_toliq_yarat','topik_shablon','topik_import','shablon_yukla','shablon_import','_mavzularni_parse','_yuklab_olish_sarlavhasi','_curriculum_scope','_curriculum_admin_guard','_dts_matn_normalize','_dts_sinf_normalize','_dts_chorak_normalize','_dts_fan_kodi_ol','_dts_bob_kodi_ol','_dts_bolim_kodi_ol','_dts_mavzu_kodi_ol','_dts_kichik_kodi_ol','_dts_qator_kiritish','_talaba_sinfini_ochish','_talaba_sinf_matni','_sinf_talaba_mi','_test_vaqti'}
  constants={'_TALABA_SINF_REGEX','_YOSH_GURUHI','TEST_AVTO_VAQT'}
  nodes=[]
  for node in ast.parse(path.read_text()).body:
   if isinstance(node,(ast.FunctionDef,ast.AsyncFunctionDef)) and node.name in names:node.decorator_list=[];nodes.append(node)
   elif isinstance(node,ast.Assign) and any(isinstance(n,ast.Name) and n.id in constants for n in node.targets):nodes.append(node)
  cls.code=compile(ast.Module(body=nodes,type_ignores=[]),str(path),'exec')
 def setUp(self):
  responses=ModuleType('fastapi.responses');responses.StreamingResponse=Download
  api=ModuleType('fastapi');api.responses=responses
  self.modules=patch.dict(sys.modules,{'fastapi':api,'fastapi.responses':responses});self.modules.start();self.addCleanup(self.modules.stop)
  self.reset()
 def reset(self):
  self.db=DB();self.db.sql.create_function('BTRIM',1,lambda value:str(value or '').strip())
  for column,kind in [('difficulty','TEXT'),('situation','TEXT'),('question','TEXT'),('option_a','TEXT'),('option_b','TEXT'),('option_c','TEXT'),('option_d','TEXT'),('correct_answer','TEXT'),('explanation','TEXT'),('question_type','TEXT'),('is_latex','BOOLEAN'),('image_url','TEXT'),('image_file_id','TEXT'),('audio_text','TEXT'),('language','TEXT'),('life_level','INTEGER'),('age_group','TEXT'),('time_limit','INTEGER'),('maqsad','TEXT'),('rasm_malumot','BLOB'),('rasm_turi','TEXT')]:
   self.db.sql.execute(f'ALTER TABLE generated_tests ADD COLUMN {column} {kind}')
  def execute_values(cur,sql,rows,**kwargs):
   self.assertIn('INSERT INTO generated_tests',sql)
   cur.rows=[]
   for row in rows:
    statement=sql.replace('VALUES %s','VALUES ('+','.join('?' for _ in row)+')')
    cursor=self.db.sql.execute(statement,row)
    cur.rows.extend(dict(result) for result in cursor.fetchall())
   cur.rowcount=len(rows)
  ns={'re':re,'io':io,'unicodedata':unicodedata,'Optional':Optional,'_curriculum':scope,'HTTPException':HTTPException,'_db':lambda:self.db,'_admin_tekshir':lambda token:99,'_jwt_tekshir':lambda token:5,'TopikShablonSorov':SimpleNamespace,'TestShablonSorov':SimpleNamespace,'UploadFile':object,'File':lambda *a,**k:None,'psycopg2':SimpleNamespace(Binary=lambda data:data,Error=sqlite3.Error,extras=SimpleNamespace(execute_values=execute_values))}
  exec(self.code,ns);self.ns=ns
 def scope(self,id=1,**changes):
  result=self.db.add_scope(id,**changes);self.db.commit();return result
 def topic_template(self,audience,grade='1 kurs',subject='Matematika',topics='1 / To‘plamlar\n1 / Sonlar'):
  response=self.ns['topik_shablon'](SimpleNamespace(scope_id=audience['id'],sinf=grade,fan=subject,dars_turi=audience['dars_turi'],mavzular=topics),'admin')
  response.headers['Content-Disposition'].encode('latin-1')
  return openpyxl.load_workbook(io.BytesIO(response.body))
 def import_topics(self,workbook,id=1):return asyncio.run(self.ns['topik_import']('admin',Upload(workbook),id))
 def create_topics(self,audience,grade='1 kurs',subject='Matematika'):
  self.import_topics(self.topic_template(audience,grade,subject),audience['id'])
  return [row[0] for row in self.db.sql.execute('SELECT topic_code FROM dts_tree WHERE curriculum_scope_id=? AND subject_name=? ORDER BY topic_code',(audience['id'],self.ns['_dts_matn_normalize'](subject).upper()))]
 def build_test_template(self,codes,id=1,groups=None):
  groups=groups or [SimpleNamespace(diff='oson',turi='single_choice',soni=1)]
  response=self.ns['shablon_yukla'](SimpleNamespace(scope_id=id,topic_codes=codes,guruhlar=groups,maqsad='oddiy'),'admin')
  return openpyxl.load_workbook(io.BytesIO(response.body))
 def fill(self,workbook,only_first=False):
  for sheet in discover_test_worksheets(workbook)[0]:
   for row in range(2,sheet.worksheet.max_row+1):
    if only_first and row>2:continue
    for column,value in [(4,f'{sheet.name} {row}: 1−1 nechaga teng?'),(5,0),(6,1),(7,2),(8,3),(9,'A')]:sheet.worksheet.cell(row,column,value)
  return workbook
 def import_tests(self,workbook,grade='1 kurs',subject='__all__',id=1):
  return asyncio.run(self.ns['shablon_import']('admin',Upload(workbook),grade,subject,id))
 def rows(self):return [dict(row) for row in self.db.sql.execute('SELECT * FROM generated_tests ORDER BY id')]
 def test_topic_template_round_trip_all_sections_and_lessons(self):
  cases=[('maktab','7',''),('bogcha','5-6 yosh',''),('markaz','a1','')]+[('universitet','1 kurs',lesson) for lesson in scope.LESSONS]
  for kind,grade,lesson in cases:
   with self.subTest(kind=kind,lesson=lesson):
    self.reset();audience=self.scope(institution_type=kind,dars_turi=lesson)
    wb=self.topic_template(audience,grade,subject='Boshlang‘ich matematika')
    wb['DTS_SHABLON'].cell(2,5,'1-bob');wb['DTS_SHABLON'].cell(2,6,'1-bo‘lim')
    result=self.import_topics(wb);self.assertEqual(result['added'],2)
    again=self.import_topics(wb);self.assertEqual(again['added'],0);self.assertEqual(again['mavjud'],2)
    rows=[dict(row) for row in self.db.sql.execute('SELECT * FROM dts_tree')]
    self.assertTrue(all(row['curriculum_scope_id']==1 and row['grade']==grade for row in rows))
    self.assertTrue(all((row['dars_turi'] or '')==lesson for row in rows))
 def test_test_template_round_trip_all_sections_and_lessons(self):
  cases=[('maktab','7',''),('bogcha','5-6 yosh',''),('markaz','a1','')]+[('universitet','1 kurs',lesson) for lesson in scope.LESSONS]
  for kind,grade,lesson in cases:
   with self.subTest(kind=kind,lesson=lesson):
    self.reset();audience=self.scope(institution_type=kind,dars_turi=lesson)
    codes=self.create_topics(audience,grade);wb=self.fill(self.build_test_template(codes))
    result=self.import_tests(wb,grade);self.assertEqual(result['saved'],2);self.assertEqual(len(self.rows()),2)
    self.assertEqual({row['topic_code'] for row in self.rows()},set(codes))
 def test_numeric_zero_options_and_written_answer_are_preserved(self):
  audience=self.scope();codes=self.create_topics(audience);wb=self.fill(self.build_test_template(codes))
  wb['TESTLAR'].cell(3,11,'write_answer');wb['TESTLAR'].cell(3,9,0)
  self.import_tests(wb);rows=self.rows()
  self.assertEqual(rows[0]['option_a'],'0');self.assertEqual(rows[1]['correct_answer'],'0')
  self.assertFalse(rows[0]['is_latex'])
 def test_long_subject_names_round_trip_in_multi_subject_file(self):
  audience=self.scope()
  codes=self.create_topics(audience,subject='Boshlang‘ich matematika kursi nazariyasi va metodikasi')
  codes+=self.create_topics(audience,subject='Boshlang‘ich matematika kursi amaliyoti va takrorlash')
  result=self.import_tests(self.fill(self.build_test_template(codes)))
  self.assertEqual(result['saved'],4);self.assertEqual(result['import_qilingan_varaq_soni'],2)
 def test_selecting_one_long_subject_does_not_import_another_sheet(self):
  audience=self.scope();subject='Boshlang‘ich matematika kursi nazariyasi va metodikasi'
  own=self.create_topics(audience,subject=subject);other=self.create_topics(audience,subject='Boshlang‘ich matematika kursi amaliyoti va takrorlash')
  self.db.sql.execute("INSERT INTO generated_tests(topic_code,question) VALUES(?,'Saqlansin')",(other[0],));self.db.commit()
  result=self.import_tests(self.fill(self.build_test_template(own+other)),subject=subject)
  self.assertEqual(result['saved'],2);self.assertEqual(len(self.rows()),3)
  self.assertTrue(any(row['question']=='Saqlansin' for row in self.rows()))
 def test_other_lesson_workbook_is_rejected_before_replacing_tests(self):
  audience=self.scope();self.scope(2,dars_turi='amaliy');codes=self.create_topics(audience)
  wb=self.fill(self.build_test_template(codes))
  self.db.sql.execute("INSERT INTO generated_tests(topic_code,question) VALUES(?,'Saqlansin')",(codes[0],));self.db.commit()
  with self.assertRaises(HTTPException) as error:self.import_tests(wb,id=2)
  self.assertEqual(error.exception.status_code,400);self.assertEqual(self.rows()[0]['question'],'Saqlansin')
 def test_image_ids_do_not_repeat_across_difficulty_groups(self):
  audience=self.scope();codes=self.create_topics(audience)
  groups=[SimpleNamespace(diff=diff,turi='single_choice',soni=2) for diff in ('oson','qiyin')]
  wb=self.build_test_template(codes,groups=groups)
  ids=[row[0] for row in wb['RASM_MALUMOTI'].iter_rows(min_row=2,values_only=True)]
  self.assertEqual(len(ids),8);self.assertEqual(len(set(ids)),8)
 def test_embedded_question_image_survives_import(self):
  audience=self.scope();codes=self.create_topics(audience);wb=self.fill(self.build_test_template(codes))
  png=io.BytesIO();Image.new('RGB',(8,8),'red').save(png,format='PNG');png.seek(0)
  wb['TESTLAR'].add_image(ExcelImage(png),'M2')
  result=self.import_tests(wb);self.assertEqual(result['rasm_biriktirildi'],1)
  image_row=next(row for row in self.rows() if row['rasm_malumot'])
  self.assertTrue(image_row['rasm_malumot'].startswith(b'\x89PNG'));self.assertEqual(image_row['image_url'],f"/api/test_rasmi/{image_row['id']}")
 def test_incomplete_question_does_not_erase_existing_tests(self):
  audience=self.scope();codes=self.create_topics(audience);wb=self.fill(self.build_test_template(codes))
  wb['TESTLAR'].cell(3,9).value=None
  self.db.sql.execute("INSERT INTO generated_tests(topic_code,question) VALUES(?,'Saqlansin')",(codes[0],));self.db.commit()
  with self.assertRaises(HTTPException):self.import_tests(wb)
  self.assertEqual([row['question'] for row in self.rows()],['Saqlansin'])
 def test_empty_topic_placeholder_keeps_that_topics_old_tests(self):
  audience=self.scope();codes=self.create_topics(audience);wb=self.fill(self.build_test_template(codes),only_first=True)
  self.db.sql.execute("INSERT INTO generated_tests(topic_code,question) VALUES(?,'Saqlansin')",(codes[1],));self.db.commit()
  result=self.import_tests(wb);self.assertEqual(result['saved'],1)
  self.assertEqual(len(self.rows()),2);self.assertTrue(any(row['question']=='Saqlansin' for row in self.rows()))
 def test_template_uses_selected_program_language(self):
  audience=self.scope(talim_tili='ru');wb=self.build_test_template(self.create_topics(audience))
  self.assertEqual(wb['TESTLAR'].cell(2,15).value,'ru')


 def test_visible_parent_template_has_15_or_60_blocks_not_double(self):
  for count in (15,60):
   with self.subTest(count=count):
    self.reset();audience=self.scope();cur=self.db.cursor();codes=[]
    for i in range(count):
     for leaf in ('Birinchi','Ikkinchi'):
      code,_=self.ns['_dts_qator_kiritish'](cur,'1 kurs','Matematika','1','','',f'Mavzu {i+1}',leaf,'maruza',1);codes.append(code)
    self.db.commit()
    visible=self.ns['topik_royxat']('1 kurs','Matematika','admin',scope_id=1)
    self.assertEqual(len(visible['mavzular']),count)
    wb=self.build_test_template(codes+codes)
    self.assertEqual(wb['TESTLAR'].max_row-1,count)
    self.assertEqual(wb['MALUMOT'].max_row-1,count)
    self.assertEqual(self.import_tests(self.fill(wb))['saved'],count)
    self.assertEqual(len(self.rows()),count)
    self.import_tests(wb);self.assertEqual(len(self.rows()),count)
 def test_question_count_is_per_parent_per_selected_group(self):
  audience=self.scope();codes=[];cur=self.db.cursor()
  for i in range(15):
   for leaf in ('A','B'):
    code,_=self.ns['_dts_qator_kiritish'](cur,'1 kurs','Matematika','1','','',f'Mavzu {i}',leaf,'maruza',1);codes.append(code)
  self.db.commit();groups=[SimpleNamespace(diff='oson',turi='single_choice',soni=4)]
  wb=self.build_test_template(codes,groups=groups);self.assertEqual(wb['TESTLAR'].max_row-1,60)
 def test_repeat_topic_lines_create_one_row(self):
  audience=self.scope();wb=self.topic_template(audience,topics='1 / Kirish\n1 / Kirish\n2 / Kirish\n2 / Kirish')
  self.assertEqual(wb['DTS_SHABLON'].max_row-1,2)
  self.assertEqual(self.import_topics(wb)['added'],2)
  self.assertEqual(self.import_topics(wb)['added'],0)
 def test_later_chapter_completion_reuses_topic_code_and_tests(self):
  audience=self.scope();wb=self.topic_template(audience,topics='1 / Kirish');self.import_topics(wb)
  code=self.db.sql.execute('SELECT topic_code FROM dts_tree').fetchone()[0]
  self.db.sql.execute("INSERT INTO generated_tests(topic_code,question) VALUES(?,'Old question')",(code,));self.db.commit()
  wb['DTS_SHABLON'].cell(2,5,'1-bob');wb['DTS_SHABLON'].cell(2,6,'1-bo‘lim')
  self.assertEqual(self.import_topics(wb)['added'],0)
  rows=list(self.db.sql.execute('SELECT topic_code,bob_name,bolim_name FROM dts_tree'))
  self.assertEqual(len(rows),1);self.assertEqual(rows[0][0],code);self.assertTrue(rows[0][1]);self.assertEqual(len(self.rows()),1)
 def test_all_four_year_pairs_round_trip_without_merging_same_title(self):
  for course in range(1,5):
   with self.subTest(course=course):
    self.reset();first,second=scope.semester_pair(course);audience=self.scope(kurs=course,semestr=first);grade=f'{course} kurs'
    wb=self.topic_template(audience,grade,topics=f'{first} / Kirish\n{second} / Kirish')
    self.assertEqual(wb['DTS_SHABLON'].cell(1,4).value,'Semestr')
    self.assertEqual(self.import_topics(wb)['added'],2)
    codes=[r[0] for r in self.db.sql.execute('SELECT topic_code FROM dts_tree')]
    tests=self.fill(self.build_test_template(codes));self.assertEqual(tests['TESTLAR'].max_row-1,2)
    self.assertEqual({r[10] for r in tests['MALUMOT'].iter_rows(min_row=2,values_only=True)},{first,second})
    self.assertEqual(self.import_tests(tests,grade)['saved'],2);self.import_tests(tests,grade);self.assertEqual(len(self.rows()),2)
    self.assertEqual(self.import_topics(wb)['added'],0)
 def test_foreign_period_rolls_back_entire_topic_workbook(self):
  audience=self.scope();wb=self.topic_template(audience);wb['DTS_SHABLON'].cell(3,4,3)
  with self.assertRaises(HTTPException):self.import_topics(wb)
  self.assertEqual(self.db.sql.execute('SELECT COUNT(*) FROM dts_tree').fetchone()[0],0)
 def test_test_replacement_clears_parent_aliases_only_in_same_semester(self):
  audience=self.scope();second=self.scope(2,semestr=2);cur=self.db.cursor();codes=[]
  for id in (1,2):
   for leaf in ('A','B'):
    code,_=self.ns['_dts_qator_kiritish'](cur,'1 kurs','Matematika',str(id),'','','Kirish',leaf,'maruza',id);codes.append(code)
    self.db.sql.execute("INSERT INTO generated_tests(topic_code,question) VALUES(?,?)",(code,f'Original {id}'))
  self.db.commit();wb=self.fill(self.build_test_template(codes[:2]));self.import_tests(wb)
  self.assertEqual(len(self.rows()),3)
  self.assertEqual(sum(r['question']=='Original 2' for r in self.rows()),2)
  self.assertFalse(any(r['question']=='Original 1' for r in self.rows()))
 def test_legacy_chorak_import_preserves_explicit_semester(self):
  audience=self.scope();self.scope(2,semestr=2)
  wb=openpyxl.Workbook();ws=wb.active;ws.append(['Sinf','Fan','Dars turi','Chorak','Bob','Bolim','Mavzu','Kichik mavzu','scope_id']);ws.append(['1 kurs','Matematika','maruza',1,'','','Kirish','',2])
  self.import_topics(wb)
  self.assertEqual(self.db.sql.execute('SELECT curriculum_scope_id FROM dts_tree').fetchone()[0],2)
 def test_official_program_renaming_does_not_fork_topic(self):
  audience=self.scope();self.import_topics(self.topic_template(audience,topics='1 / Kirish'))
  self.db.sql.execute("UPDATE curriculum_scopes SET yonalish_nomi='Renamed',yonalish_key='renamed' WHERE id=1");self.db.commit()
  self.assertEqual(self.import_topics(self.topic_template(audience,topics='1 / Kirish'))['added'],0)


 def test_explanations_with_same_title_require_semester_and_remain_separate(self):
  audience=self.scope();self.import_topics(self.topic_template(audience,topics='1 / Kirish\n2 / Kirish'))
  self.db.sql.execute("CREATE TABLE curriculum_explanations(curriculum_scope_id INTEGER,sinf TEXT,fan TEXT,mavzu_nomi TEXT,tushuntirish TEXT,yaratilgan_at TEXT,UNIQUE(curriculum_scope_id,sinf,fan,mavzu_nomi))");self.db.commit()
  wb=openpyxl.Workbook();ws=wb.active;ws.append(['Sinf','Fan','Mavzu','Tushuntirish']);ws.append(['1 kurs','Matematika','Kirish','First'])
  with self.assertRaises(HTTPException):asyncio.run(self.ns['tushuntirish_import']('admin',Upload(wb),1))
  self.assertEqual(self.db.sql.execute('SELECT COUNT(*) FROM curriculum_explanations').fetchone()[0],0)
  ws.cell(1,5,'Semestr');ws.cell(2,5,1);ws.append(['1 kurs','Matematika','Kirish','Second',2])
  result=asyncio.run(self.ns['tushuntirish_import']('admin',Upload(wb),1));self.assertEqual(result['saqlandi'],2)
  codes=list(self.db.sql.execute('SELECT d.topic_code,cs.semestr FROM dts_tree d JOIN curriculum_scopes cs ON cs.id=d.curriculum_scope_id'))
  for code,semester in codes:
   result=self.ns['mavzu_tushuntirish_ol']('1 kurs','Matematika','kirish','student',code)
   self.assertEqual(result['tushuntirish'],'First' if semester==1 else 'Second')
  self.assertFalse(self.ns['mavzu_tushuntirish_ol']('1 kurs','Matematika','kirish','student')['topildi'])

if __name__=='__main__':unittest.main()

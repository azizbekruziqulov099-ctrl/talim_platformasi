"""REV152: xavfsizlik — fayl tekshiruvi, cheklovlar, Telegram bot tekshiruvi."""
import io
import os
import unittest
import zipfile

from fastapi import HTTPException

from platform_core import limitlar, xavfsiz_fayl as X


def _docx(qoshimcha=None) -> bytes:
    import docx
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph("Salom, bu sinov hujjati.")
    d.save(buf)
    if not qoshimcha:
        return buf.getvalue()
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in src.infolist():
            z.writestr(i, src.read(i))
        for nom, data in qoshimcha.items():
            z.writestr(nom, data)
    return out.getvalue()


def _xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Ism"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _png(fmt="PNG") -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 50, 50)).save(buf, fmt)
    return buf.getvalue()


def _pdf() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (100, 100), (255, 255, 255)).save(buf, "PDF")
    return buf.getvalue()


JS_PDF = (b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R/OpenAction 3 0 R>>endobj\n"
          b"2 0 obj<</Type/Pages/Kids[4 0 R]/Count 1>>endobj\n"
          b"3 0 obj<</S/JavaScript/JS(app.alert('x'))>>endobj\n"
          b"4 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 100 100]>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF")


class FaylTekshiruv(unittest.TestCase):
    def rad(self, data, nom, matn=None, ruxsat=X.KUTUBXONA):
        with self.assertRaises(HTTPException) as c:
            X.tekshir(data, nom, ruxsat)
        if matn:
            self.assertIn(matn, str(c.exception.detail))
        return c.exception

    def test_haqiqiy_fayllar_otadi_va_tur_serverda_belgilanadi(self):
        self.assertEqual(X.tekshir(_docx(), "reja.docx").mime, X.MIME["docx"])
        self.assertEqual(X.tekshir(_xlsx(), "jadval.xlsx").tur, "xlsx")
        self.assertEqual(X.tekshir(_pdf(), "kitob.pdf").mime, "application/pdf")
        self.assertEqual(X.tekshir(_png(), "rasm.png").mime, "image/png")
        self.assertEqual(X.tekshir(_png("JPEG"), "foto.JPG").tur, "jpeg")

    def test_faqat_word_excel_pdf_rasm(self):
        self.rad(b"hello", "a.txt", "Faqat")
        self.rad(b"PK\x03\x04", "slayd.pptx", "Faqat")
        self.rad(b"MZ\x90\x00", "dastur.exe", "Faqat")
        self.rad(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, "eski.doc", ".docx")

    def test_nomi_ozgartirilgan_fayl_rad_etiladi(self):
        self.rad(b"<html><script>alert(1)</script></html>", "hujjat.pdf", "mos emas")
        self.rad(b"MZ\x90\x00" + b"\0" * 100, "hujjat.docx", "mos emas")
        self.rad(_png(), "rasm.jpg", "mos emas")
        self.rad(_png(), "rasm.pdf", "mos emas")

    def test_makros_va_obyektlar_rad_etiladi(self):
        self.rad(_docx({"word/vbaProject.bin": b"x" * 100}), "makros.docx", "makros")
        self.rad(_docx({"word/embeddings/oleObject1.bin": b"x" * 100}), "obj.docx", "joylangan")
        self.rad(_docx({"word/activeX/activeX1.xml": b"<x/>"}), "ax.docx", "ActiveX")
        # diagramma uchun joylangan Excel — mumkin
        X.tekshir(_docx({"word/embeddings/Microsoft_Excel_Worksheet.xlsx": _xlsx()}), "diagramma.docx")

    def test_zip_bomba_va_yol_chiqishi(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            src = zipfile.ZipFile(io.BytesIO(_docx()))
            for i in src.infolist():
                z.writestr(i, src.read(i))
            with z.open("word/media/katta.bin", "w") as f:
                bolak = b"\0" * (8 * 1024 * 1024)
                for _ in range(34):
                    f.write(bolak)
        self.rad(out.getvalue(), "bomba.docx", "zip-bomba")
        self.rad(_docx({"../../etc/x.xml": b"<x/>"}), "yol.docx", "shubhali")

    def test_soxta_word_tuzilmasi(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            z.writestr("index.html", "<script>alert(1)</script>")
        self.rad(out.getvalue(), "soxta.docx", "haqiqiy Word")

    def test_xavfli_pdf(self):
        self.rad(JS_PDF, "skript.pdf", "PDF")
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter()
        for p in PdfReader(io.BytesIO(_pdf())).pages:
            w.add_page(p)
        w.encrypt("123")
        buf = io.BytesIO()
        w.write(buf)
        self.rad(buf.getvalue(), "parolli.pdf", "Parol")

    def test_bosh_va_katta(self):
        self.rad(b"", "a.pdf", "bo‘sh")
        e = self.rad(_pdf(), "a.pdf", "katta", ruxsat=X.KUTUBXONA) if False else None
        with self.assertRaises(HTTPException) as c:
            X.tekshir(_pdf(), "a.pdf", X.KUTUBXONA, max_bayt=100)
        self.assertEqual(c.exception.status_code, 413)

    def test_xavfsiz_qaytarish(self):
        r = X.javob(b"<html>", "x.html")
        self.assertIn("attachment", r.headers["content-disposition"])
        self.assertEqual(r.headers["content-type"], "application/octet-stream")
        self.assertIn("sandbox", r.headers["content-security-policy"])
        r = X.javob(_pdf(), "kitob.pdf")
        self.assertIn("inline", r.headers["content-disposition"])
        self.assertEqual(r.headers["x-content-type-options"], "nosniff")

    def test_chat_media_va_eski_html(self):
        webm = b"\x1a\x45\xdf\xa3" + b"\0" * 50
        self.assertEqual(X.media_turi(webm, "audio/webm;codecs=opus"), "audio/webm")
        with self.assertRaises(HTTPException):
            X.media_turi(b"<html><script>", "audio/webm")
        with self.assertRaises(HTTPException):
            X.media_turi(webm, "text/html")
        # bazada eski yozuv: text/html turi bilan saqlangan «audio» — brauzerda ochilmaydi
        r = X.chat_javob(b"<html><script>alert(1)</script>", "x.webm", "text/html")
        self.assertIn("attachment", r.headers["content-disposition"])
        self.assertNotIn("html", r.headers["content-type"])


class Cheklovlar(unittest.TestCase):
    def test_xotiradagi_hisoblagich(self):
        limitlar.ulash(None)
        for _ in range(3):
            limitlar.sarfla("test", "a", 1, 3, 60)
        with self.assertRaises(HTTPException) as c:
            limitlar.sarfla("test", "a", 1, 3, 60)
        self.assertEqual(c.exception.status_code, 429)
        limitlar.sarfla("test", "b", 1, 3, 60)   # boshqa subyekt — alohida
        self.assertEqual(limitlar.hisobla("belgi", "x", 1000, 60) + limitlar.hisobla("belgi", "x", 500, 60), 1000 + 1500)

    def test_muhit_soni(self):
        os.environ["TEST_SON"] = "7"
        self.assertEqual(limitlar.son("TEST_SON", 3), 7)
        os.environ["TEST_SON"] = "abc"
        self.assertEqual(limitlar.son("TEST_SON", 3), 3)


class TelegramBot(unittest.TestCase):
    def test_begona_bot_qabul_qilinmaydi(self):
        import kabutar_auth as A
        svc = A.AuthService.__new__(A.AuthService)
        svc.bot_username = ""
        svc._bot_token = ""
        svc._username_checked_at = 0.0
        svc._ensure_bot_username = lambda: None
        eski = A.bot_username_from_token
        A.bot_username_from_token = lambda t: "hacker_bot"
        try:
            begona = "123456789:" + "A" * 35
            self.assertFalse(svc.verify_bot_token(begona), "nom noma'lum bo'lsa — rad")
            self.assertEqual(svc.bot_username, "", "begona bot nomi o'zlashtirilmaydi")
            svc.bot_username = "kabutar_bot"
            self.assertFalse(svc.verify_bot_token(begona))
            A.bot_username_from_token = lambda t: "kabutar_bot"
            svc.__dict__.pop("_bot_token_cache", None)
            self.assertTrue(svc.verify_bot_token("987654321:" + "B" * 35))
        finally:
            A.bot_username_from_token = eski


if __name__ == "__main__":
    unittest.main()

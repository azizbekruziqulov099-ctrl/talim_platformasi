import io
import unittest

import openpyxl

from modules.ai_miya_oddiy import COLUMNS, SAMPLE_ROWS, parse_simple, split_solution, template_workbook
from modules.dars_export import lesson_docx, lesson_pdf, plain
from modules.dars_xonasi import build_lesson

KEYS = [c[0] for c in COLUMNS]


def filled(rows, book=("Matematika 5", "Matematika", 5), image_at=None):
    wb = template_workbook()
    k = wb["KITOB"]
    for r in range(1, 15):
        label = k.cell(r, 1).value
        if label == "Kitob nomi *": k.cell(r, 2).value = book[0]
        if label == "Fan *": k.cell(r, 2).value = book[1]
        if label == "Sinf *": k.cell(r, 2).value = book[2]
    ws = wb["MAVZULAR"]
    start = len(SAMPLE_ROWS) + 3
    for i, row in enumerate(rows):
        for j, key in enumerate(KEYS, 1):
            ws.cell(start + i, j, row.get(key, ""))
    if image_at:
        from openpyxl.drawing.image import Image as XLImage
        from PIL import Image
        buf = io.BytesIO(); Image.new("RGB", (40, 30), (200, 90, 40)).save(buf, "PNG"); buf.seek(0)
        col = KEYS.index(image_at[1]) + 1
        ws.add_image(XLImage(buf), f"{openpyxl.utils.get_column_letter(col)}{start + image_at[0]}")
    out = io.BytesIO(); wb.save(out)
    return openpyxl.load_workbook(io.BytesIO(out.getvalue()))


def own(code="5-01"):
    rows = [dict(r) for r in SAMPLE_ROWS]
    rows[0]["mavzu_kodi"] = code
    return rows


def units_of(p):
    return [{"unit_code": r[k], "unit_kind": kind, "title": "", "body": "", "payload": r}
            for sheet, k, kind in (("11_DARS_SSENARIY", "step_id", "lesson_step"), ("12_TUSHUNMADIM", "variant_id", "variant"))
            for r in p[sheet]]


class OddiyShablonTests(unittest.TestCase):
    def test_template_is_two_sheets_one_row_per_concept_without_tests(self):
        wb = template_workbook()
        self.assertEqual(wb.sheetnames, ["KITOB", "MAVZULAR"])
        headers = [c.value for c in wb["MAVZULAR"][2]]
        self.assertFalse(any("test" in str(h).lower() for h in headers))
        parsed = parse_simple(filled([]))
        self.assertTrue(any("Birorta tushuncha" in e["message"] for e in parsed["errors"]))

    def test_each_row_becomes_its_own_concept_with_variants_examples_and_task(self):
        parsed = parse_simple(filled(own(), image_at=(0, "rasm")), {"kataklar_3_8.png"})
        self.assertEqual(parsed["errors"], [])
        p = parsed["payload"]
        self.assertEqual(len(p["02_DTS_XARITA"]), 1)
        self.assertEqual(p["06_MASHQLAR"], [])
        steps = p["11_DARS_SSENARIY"]
        kinds = [s["qadam_turi"] for s in steps]
        self.assertEqual(kinds[:7], ["kirish", "qoida", "misol", "misol", "misol", "misol", "birga"])
        self.assertEqual(kinds.count("qoida"), 2)  # ikkita tushuncha = ikkita qator
        first_rule = steps[1]
        self.assertTrue(all(v["step_id"] == first_rule["step_id"] for v in p["12_TUSHUNMADIM"][:5]))
        second_rule = [s for s in steps if s["qadam_turi"] == "qoida"][1]
        self.assertTrue(any(v["step_id"] == second_rule["step_id"] for v in p["12_TUSHUNMADIM"]))
        self.assertEqual(len({s["sahna"] for s in steps[2:6]}), 1)  # misol bitta doskada
        self.assertEqual(steps[3]["doska_matni"], "Maxraj = 8")
        self.assertIn("sakkiz", steps[3]["ovoz_matni"])
        self.assertEqual(first_rule["media_id"], "pitsa_8.png")
        self.assertIn("pitsa_8.png", parsed["media"])
        self.assertEqual(len(parsed["media"]), 1)
        lesson = build_lesson({"topic_code": "5-01", "mavzu": "Oddiy kasr"}, units_of(p), lambda m: None,
                              [{"id": "t", "savol": "Q", "variantlar": list("abcd"), "togri": 1, "izoh": ""}])
        self.assertEqual(len(lesson["savollar"]), 1)  # testlar bazadan
        self.assertTrue(lesson_pdf(lesson, {}).startswith(b"%PDF"))
        self.assertEqual(lesson_docx(lesson, {})[:2], b"PK")

    def test_clear_errors_point_to_row_and_column(self):
        rows = own()
        rows.append({"mavzu_kodi": "", "tushuncha": "Bo'sh", "masala_javob": "5"})
        rows.append({"tushuncha": "Katta", "tushuntirish": "x" * 1600})
        parsed = parse_simple(filled(rows), {"pitsa_8.png", "kataklar_3_8.png"})
        messages = " | ".join(e["message"] for e in parsed["errors"])
        self.assertIn("Tushuntirish yoki doska", messages)
        self.assertIn("1500", messages)
        self.assertTrue(all(e["sheet"] == "MAVZULAR" and e["row"] > 2 for e in parsed["errors"]))

    def test_missing_named_picture_is_a_warning(self):
        parsed = parse_simple(filled(own()))
        self.assertTrue(any("pitsa_8.png" in w["message"] for w in parsed["warnings"]))

    def test_solution_lines_and_formula_text(self):
        self.assertEqual(split_solution("2+3 || ikki qo'shuv uch\n=5"), [("2+3", "ikki qo'shuv uch"), ("=5", "=5")])
        self.assertEqual(plain(r"$\frac{\text{olingan}}{\text{jami}}$"), "olingan/jami")


if __name__ == "__main__":
    unittest.main()

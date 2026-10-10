"""REV80/REV87: bog'cha kitobi (dastur JSON) → platformaga yuklanadigan fayllar.

Chiqadi (har yosh guruhi uchun):
  1) <fan>_<yosh>_1_mavzular.xlsx — Mavzular importi (Admin → Mavzular → Import, dastur: «Bog'cha — umumiy katalog»)
  2) <fan>_<yosh>_2_ai_miya.xlsx  — AI miya kitobi (Admin → Shablon → Kitob darslari). Tayyor rasmlar
     (--rasmlar papkasidan, fayl nomi bo'yicha) «Rasm» katagiga o'zi joylanadi.
Hamma yosh uchun bitta:
  3) <fan>_rasmlar_royxati.xlsx — kerakli rasmlar: fayl nomi + GPT buyrug'i + holati (bor / kerak)

Dastur yonida `<nom>_enrich.json` bo'lsa (masalan en_56_enrich.json) — u avtomatik qo'shiladi (inglizcha misol
gaplar, kirishlar, hayotiy vaziyat darslari). Bir fandagi kitoblar yosh tartibida beriladi: har kitob oldingi yosh
kitobidan «O'tgan yilni eslaymiz» darsini oladi.

Ishlatish:  python tools/bogcha_kitob.py <chiqish_papka> [--rasmlar <papka>] dastur1.json dastur2.json ...
Eski usul:  python tools/bogcha_kitob.py <chiqish_papka> "Ingliz tili" EN kitob1.json ...   (tayyor kitob JSON)
"""
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from modules.ai_miya_varoq import template_workbook  # noqa: E402
from tools.bogcha_teg import fix_book  # noqa: E402

STYLE = "cute cartoon, soft flat colors, white background, no text, square 1024x1024"
IMAGE_PX = 110


def age_short(age):
    return "".join(ch for ch in age if ch.isdigit())[:2]


def topics_workbook(book, fan):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mavzular"
    ws.append(["Sinf", "Fan", "Chorak", "Bob", "Bo'lim", "Mavzu", "Kichik mavzu"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B57")
    for topic in book["topics"]:
        ws.append([book["age"], fan, 1, book.get("book_title", fan), "", topic["name"], ""])
    for col, width in zip("ABCDEFG", (12, 16, 8, 30, 10, 36, 14)):
        ws.column_dimensions[col].width = width
    return wb


def ai_workbook(book, fan, prefix, til="uz"):
    prefill = []
    for topic in book["topics"]:
        rows = []
        for row in topic["rows"]:
            rows.append({
                "turi": row.get("turi", ""), "sarlavha": row.get("sarlavha", ""), "matn": row.get("matn", ""),
                "doska": row.get("doska", ""), "rasm": row.get("rasm") or "", "variantlar": row.get("variantlar", ""),
                "javob": row.get("javob", ""),
                "yechim": row.get("yechim") or ("Barakalla! Juda yaxshi bajarding! 🌟" if row.get("turi") == "topshiriq" else ""),
                "sodda": row.get("sodda", ""), "boshqa_usul": row.get("boshqa_usul", ""),
            })
        prefill.append({"mavzu_kodi": "", "mavzu_nomi": topic["name"], "mavzu_raqami": topic["no"], "daraja": 1, "rows": rows})
    meta = {"kitob_nomi": book.get("book_title") or f"{fan} {book['age']}", "fan": fan, "sinf": book["age"],
            "til": til, "kod_prefiksi": f"{prefix}{age_short(book['age'])}", "mualliflar": "Kabutar Ta'lim"}
    return template_workbook(prefill, meta, blank_topics=0)


def embed_images(wb, image_dir):
    """«Rasm» katagidagi fayl nomi bo'yicha rasmni katakka joylaydi. Topilmagan nom tozalanadi (import ogohlantirmasin).
    Qaytaradi: (joylangan soni, topilmagan nomlar to'plami)."""
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage
    image_dir = Path(image_dir) if image_dir else None
    placed, missing, cache, used = 0, set(), {}, set()
    embed_images.used = used
    for ws in wb.worksheets:
        if ws.title in ("KITOB", "NAMUNA"):
            continue
        col = header_row = None
        for r in range(1, 8):
            for c in range(1, ws.max_column + 1):
                if str(ws.cell(r, c).value or "").strip() == "Rasm":
                    col, header_row = c, r
        if not col:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            name = str(ws.cell(r, col).value or "").strip()
            if not name:
                continue
            path = image_dir / name if image_dir else None
            svg = image_dir / (Path(name).stem + ".svg") if image_dir else None
            if (not path or not path.is_file()) and svg and svg.is_file():
                # REV99: jonli SVG rasm — katakda .svg nomi qoladi (ZIP ichida yuboriladi), Excel'da ko'rinishi uchun
                # png_preview/<nom>.png bo'lsa o'sha joylanadi.
                ws.cell(r, col).value = svg.name
                used.add(svg.name)
                name, path = svg.name, image_dir.parent / "png_preview" / (svg.stem + ".png")
                if not path.is_file():
                    continue
            elif not path or not path.is_file():
                missing.add(name)
                ws.cell(r, col).value = None
                continue
            if name in cache:
                continue   # shu kitobda allaqachon joylangan — fayl nomi qoladi, import o'sha rasmni ishlatadi
            if name not in cache:
                with PILImage.open(path) as im:
                    im = im.convert("RGB")
                    im.thumbnail((256, 256))
                    buf = io.BytesIO()
                    im.save(buf, "JPEG", quality=82)
                cache[name] = buf.getvalue()
            pic = XLImage(io.BytesIO(cache[name]))
            pic.width = pic.height = IMAGE_PX
            ws.add_image(pic, ws.cell(r, col).coordinate)
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, IMAGE_PX * 0.78)
            placed += 1
    return placed, missing


def images_workbook(books, fan, image_dir=None):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    image_dir = Path(image_dir) if image_dir else None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rasmlar"
    ws.append(["Yosh", "Mavzu", "Fayl nomi", "Nima", "Emoji (hozircha)", "GPT uchun buyruq (inglizcha)", "Holati"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B57")
    seen = set()
    rows = []
    for book in books:
        for topic in book["topics"]:
            if topic.get("image_scene") and topic["image_scene"] not in seen:
                seen.add(topic["image_scene"])
                rows.append([book["age"], topic["name"], topic["image_scene"], "Hayotiy vaziyat sahnasi", "",
                             topic.get("image_scene_prompt") or f"A cheerful scene for the lesson '{topic['name']}', {STYLE}"])
            for word in topic.get("words", []):
                if not word.get("image") or word["image"] in seen:
                    continue  # takror darslaridagi so'zlar ikkinchi marta yozilmaydi
                seen.add(word["image"])
                prompt = word.get("image_prompt") or f"{word.get('en')}, {STYLE}"
                rows.append([book["age"], topic["name"], word.get("image", ""), f"{word.get('en', '')} — {word.get('uz', '')}",
                             word.get("emoji", ""), prompt])
    have = lambda name: bool(image_dir and ((image_dir / name).is_file() or (image_dir / (Path(name).stem + ".svg")).is_file()))  # noqa: E731
    rows.sort(key=lambda r: have(r[2]))   # kerak bo'lganlari tepada
    missing = 0
    for r in rows:
        ok = have(r[2])
        missing += not ok
        ws.append(r + ["✅ bor" if ok else "🎨 kerak"])
        if not ok:
            ws.cell(ws.max_row, 7).font = Font(bold=True, color="B5541C")
    for col, width in zip("ABCDEFG", (10, 26, 30, 26, 10, 90, 10)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return wb, missing, len(rows)


def build_books(paths, izoh=None):
    """Dastur JSON'lar (yosh tartibida) → kitoblar. Enrich va oldingi yosh avtomatik.
    izoh: tushuntirish tili (uz — asl; ru | en — bogcha_content/izoh/<izoh>.json lug'ati bilan)."""
    from tools.bogcha_spiral import build_book
    books, prev_by_subject = [], {}
    for path in paths:
        path = Path(path)
        data = json.loads(path.read_text(encoding="utf-8"))
        if "topics" in data:        # allaqachon tayyor kitob
            books.append(data)
            continue
        enrich_path = path.with_name(f"{path.stem}_enrich.json")
        enrich = json.loads(enrich_path.read_text(encoding="utf-8")) if enrich_path.is_file() else None
        subject = data.get("subject")
        if izoh and izoh != "uz":
            data["izoh"] = izoh
        book = build_book(data, enrich, prev_by_subject.get(subject))
        prev_by_subject[subject] = data
        books.append(book)
    return books


IZOH_NOMI = {"ru": "rus", "en": "ingliz"}


def write_groups(out, groups, prefix=None, image_dir=None, izoh="uz"):
    """{fan: [kitoblar]} → papkalar. izoh != uz: papka/fayl nomida «_izoh_<izoh>», meta «til» = izoh,
    fan nomi «Ingliz tili (izoh: rus)»."""
    out = Path(out)
    report = []
    for subject, items in groups.items():
        slug = subject.lower().replace(" ", "_").replace("'", "").replace("‘", "")
        if izoh != "uz":
            slug = f"{slug}_izoh_{izoh}"
            subject = f"{subject} (izoh: {IZOH_NOMI.get(izoh, izoh)})"
        folder = out / slug
        folder.mkdir(parents=True, exist_ok=True)
        for book in items:
            fix_book(book, izoh)   # REV111: har til o'z ovozida o'qilsin (teglar)
            a = age_short(book["age"])
            topics_workbook(book, subject).save(folder / f"{slug}_{a[0]}-{a[1]}_yosh_1_mavzular.xlsx")
            wb = ai_workbook(book, subject, prefix or book.get("prefix") or "BK", til=izoh)
            placed, missing = embed_images(wb, image_dir)
            wb.save(folder / f"{slug}_{a[0]}-{a[1]}_yosh_2_ai_miya.xlsx")
            report.append(f"{subject} {book['age']}: {len(book['topics'])} dars, rasm joylandi {placed}, rasm kerak {len(missing)}")
        wb, missing, total = images_workbook(items, subject, image_dir)
        wb.save(folder / f"{slug}_rasmlar_royxati.xlsx")
        report.append(f"{subject}: rasmlar ro'yxati — jami {total}, chizish kerak {missing}")
    return report


def main(out, *args):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    args = list(args)
    image_dir = None
    if "--rasmlar" in args:
        i = args.index("--rasmlar")
        image_dir = args[i + 1]
        del args[i:i + 2]
    fan = prefix = None
    if args and not str(args[0]).endswith(".json"):
        fan, prefix, args = args[0], args[1], args[2:]
    books = build_books(args)
    groups = {}
    for book in books:
        subject = fan or book.get("subject") or "Fan"
        groups.setdefault(subject, []).append(book)
    report = write_groups(out, groups, prefix, image_dir)
    print("\n".join(report))
    return report


if __name__ == "__main__":
    main(*sys.argv[1:])

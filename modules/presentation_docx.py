"""Editable Word content forms for the REV50 planned slide import protocol.

This is a text form: formula slots intentionally contain source LaTeX, and image
slots describe images which the user attaches in the presentation workspace.
No external resources or new dependencies are used.
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor


# Keep in sync with frontend presentations/layouts.js field counts. Hidden
# values still receive tags so selecting another layout never loses content.
LAYOUT_FIELDS = {
    "text": (1, 0), "formula": (1, 0), "image": (1, 1), "cover": (1, 1),
    "two_columns": (2, 0), "two_images": (2, 2), "three_cards": (3, 0), "steps": (3, 0),
}
FIELD_HINTS = {
    "SARLAVHA": "Slayd sarlavhasini yozing. Ko‘pi bilan 100 belgi.",
    "MATN1": "Birinchi matnni yozing. Ko‘pi bilan 600 belgi.",
    "MATN2": "Ikkinchi matnni yozing. Ko‘pi bilan 400 belgi.",
    "MATN3": "Uchinchi matnni yozing. Ko‘pi bilan 400 belgi.",
    "FORMULA": "Ixtiyoriy formula manbasini LaTeX ko‘rinishida yozing. Ko‘pi bilan 400 belgi.",
    "MISOL": "Ixtiyoriy misol yoki savol yozing. Ko‘pi bilan 250 belgi.",
    "RASM1": "Birinchi rasm uchun tavsif yozing. Ko‘pi bilan 240 belgi. Rasmni saytda yuklang.",
    "RASM1_IZOH": "Birinchi rasm ostidagi izoh. Ko‘pi bilan 100 belgi.",
    "RASM2": "Ikkinchi rasm uchun tavsif yozing. Ko‘pi bilan 240 belgi. Rasmni saytda yuklang.",
    "RASM2_IZOH": "Ikkinchi rasm ostidagi izoh. Ko‘pi bilan 100 belgi.",
}
INSTRUCTIONS = (
    "Ushbu Word shakli tanlangan taqdimotning har bir slaydi uchun mazmun tayyorlashga yordam beradi. "
    "Har bir slaydni to‘ldiring, so‘ng faylni Taqdimotlar oynasida import qilib natijani ko‘rib chiqing.",
    "Slaydlar soni va tartibini, SLAYD, ID, MAKET hamda BOLIM qiymatlarini o‘zgartirmang. "
    "Kvadrat qavs ichidagi teglarni saqlang. Matnni tegdan keyin yozing; ixtiyoriy maydon bo‘sh qolishi mumkin.",
    "MATN1, MATN2 va MATN3 alohida matn joylari. FORMULA maydoniga LaTeX manbasi, MISOL maydoniga oddiy matn yoziladi. "
    "Teglar yonidagi belgi chegaralariga amal qiling. Tayyor javoblar o‘rniga burchak qavsli o‘rinbosarlar yozmang.",
    "RASM1 va RASM2 rasm tavsiflari xolos. Rasmlar avtomatik yaratilmaydi yoki topilmaydi. "
    "Rasmlarni saytda yuklang; oldingi rasmlar va slayd dizayni saqlanadi. Word ichidagi rasmlar import qilinmaydi.",
    "RASM1_IZOH va RASM2_IZOH rasmlar ostidagi izohlardir; ular slaydda tegishli rasm mavjud bo‘lsa ko‘rinadi. "
    "Maketingizda yashiringan, oldin yozilgan maydonlar ham mazmunni saqlash uchun shaklga kiritiladi.",
    "Yordam satrlari javob emas va import qilinmaydi. Yangi javob satrini kvadrat qavsli teg bilan yoki Yordam ko‘rsatmasi bilan boshlamang. "
    "Word bezaklari import qilinmaydi; bu fayl slayd mazmunini to‘ldirish shaklidir.",
)


def _value(value):
    value = str(value or "").strip().replace("\r\n", "\n").replace("\r", "\n")
    # Reserved line syntax cannot be losslessly represented by this text form.
    # Tell the user to amend it instead of silently changing their content.
    if any(re.match(r"^\s*\[", line) or re.fullmatch(r"\(Yordam: .*\)", line.strip()) for line in value.split("\n")):
        raise ValueError("Matn satri teg yoki Yordam ko‘rsatmasi bilan boshlanmasin; uni oddiy gap shaklida yozing")
    if re.search(r"<[^<>\n]+>", value):
        raise ValueError("Word shaklida burchak qavsli o‘rinbosar o‘rniga mazmun yozing yoki maydonni bo‘sh qoldiring")
    return value


def slide_fields(slide):
    """Ordered (tag, value) entries, exactly matching the JavaScript builder."""
    layout = slide.get("layout", "text")
    if layout not in LAYOUT_FIELDS:
        raise ValueError("Slayd maketi noto‘g‘ri")
    text_count, image_count = LAYOUT_FIELDS[layout]
    result = [("ID", slide["id"]), ("MAKET", layout), ("BOLIM", slide.get("section", "")),
              ("SARLAVHA", slide.get("title", "")), ("MATN1", slide.get("body", ""))]
    for number in (2, 3):
        value = slide.get(f"body{number}", "")
        if text_count >= number or value:
            result.append((f"MATN{number}", value))
    result.extend((("FORMULA", slide.get("formula", "")), ("MISOL", slide.get("example", ""))))
    for number, prefix in ((1, "image"), (2, "image2")):
        prompt, caption = slide.get(prefix + "_prompt", ""), slide.get(prefix + "_caption", "")
        if image_count >= number or prompt or caption:
            result.extend(((f"RASM{number}", prompt), (f"RASM{number}_IZOH", caption)))
    return [(tag, _value(value)) for tag, value in result]


def build_tagged_template(document):
    """Plain-text mirror of the Word form, useful for protocol tests."""
    paragraphs = ["Taqdimot mazmunini to‘ldirish shakli", *INSTRUCTIONS]
    for index, slide in enumerate(document["slides"], 1):
        paragraphs.append(f"[SLAYD:{index}]")
        for tag, value in slide_fields(slide):
            paragraphs.append(f"[{tag}]" + (" " + value if value else ""))
            if not value and tag in FIELD_HINTS:
                paragraphs.append(f"(Yordam: {FIELD_HINTS[tag]})")
        paragraphs.append("[/SLAYD]")
    return "\n".join(paragraphs)


def _add_field(doc, tag, value):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(f"[{tag}]").bold = True
    if value:
        paragraph.add_run(" " + value)
    elif tag in FIELD_HINTS:
        paragraph.paragraph_format.keep_with_next = True
        hint = doc.add_paragraph(f"(Yordam: {FIELD_HINTS[tag]})")
        hint.paragraph_format.space_after = Pt(4)
        for run in hint.runs:
            run.italic = True
            run.font.color.rgb = RGBColor.from_string("404040")


def build_template_docx(document):
    """Return a real OOXML DOCX with instructions and one form per slide.

    Page breaks live on each slide heading, outside the tagged block, so the
    hardened paragraph extractor sees the same text protocol as copied text.
    """
    # Validate all field syntax before authoring any output.
    planned = [slide_fields(slide) for slide in document["slides"]]
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        # The bundled default may carry a theme-colored Title underline.
        # A content form uses hierarchy and whitespace, not decorative rules.
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = False
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["Title"].paragraph_format.space_after = Pt(12)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(9)
    doc.core_properties.title = "Taqdimot mazmunini to‘ldirish shakli"
    doc.core_properties.subject = document.get("title", "")
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.add_paragraph("Taqdimot mazmunini to‘ldirish shakli", "Title")
    doc.add_paragraph(f"Taqdimot: {document.get('title', '')}")
    doc.add_paragraph(f"Slaydlar soni: {len(planned)}")
    for instruction in INSTRUCTIONS:
        doc.add_paragraph(instruction)
    for index, entries in enumerate(planned, 1):
        heading = doc.add_paragraph(f"Slayd {index}", "Heading 1")
        heading.paragraph_format.page_break_before = True
        paragraph = doc.add_paragraph(f"[SLAYD:{index}]")
        paragraph.paragraph_format.space_after = Pt(5)
        for tag, value in entries:
            _add_field(doc, tag, value)
        doc.add_paragraph("[/SLAYD]")
    # The import guard counts extracted text, including paragraph separators.
    # Never issue a content form that our own unchanged importer cannot read.
    if sum(len(paragraph.text) + 2 for paragraph in doc.paragraphs) > 100000:
        raise ValueError("Word shakli matni 100 000 belgidan oshdi. Matnni qisqartiring yoki taqdimotni bo‘ling")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()

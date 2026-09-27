"""Ochiq dars ishlanmasi: dars xonasidagi darsni PDF va Word (.docx) ga chiqaradi."""
import re
import threading
from io import BytesIO
from pathlib import Path

FONT_ROOT = Path(__file__).with_name("pdf_fonts")
_LOCK = threading.Lock()
STEP_LABELS = {
    "kirish": "Kirish (motivatsiya)", "tushuntirish": "Yangi mavzu tushuntirilishi", "qoida": "Qoida",
    "misol": "Misol yechish", "birga": "Birgalikda ishlash", "mashq": "Mashq", "xulosa": "Xulosa",
}
VARIANT_LABELS = {"sodda": "Soddaroq", "hikoya": "Hikoya orqali", "rasm": "Rasm bilan",
                  "boshqa_usul": "Boshqa usulda", "takrorlash": "Oldingi mavzuni takrorlash"}


def plain(text):
    """Doskadagi LaTeX'ni oddiy o'qiladigan matnga aylantiradi (hujjat uchun)."""
    t = str(text or "")
    t = re.sub(r"\\(?:text|mathrm|mathbf)\{([^{}]*)\}", r"\1", t)
    for _ in range(4):
        t = re.sub(r"\\[dt]?frac\{([^{}]*)\}\{([^{}]*)\}", r"\1/\2", t)
    t = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", t)
    for src, dst in (("\\cdot", "·"), ("\\times", "×"), ("\\div", ":"), ("\\le", "≤"), ("\\ge", "≥"),
                     ("\\ne", "≠"), ("\\pi", "π"), ("\\;", " "), ("\\,", " "), ("\\quad", " ")):
        t = t.replace(src, dst)
    t = t.replace("$", "").replace("{", "").replace("}", "").replace("\\", "")
    return re.sub(r"[ \t]+", " ", t).strip()


def _voice(text):
    """O'qituvchi matnidan doska belgilarini ([1], [2]) olib tashlaydi."""
    return re.sub(r"\s*\[\d{1,2}\]\s*", " ", str(text or "")).strip()


def _no_emoji(text):
    return re.sub("[\U0001F000-\U0001FAFF\u2600-\u27BF\uFE0F\u200d]", "", str(text or "")).strip()


def _png(data):
    try:
        from PIL import Image
        image = Image.open(BytesIO(data))
        image.thumbnail((1400, 1400))
        if image.mode not in ("RGB", "RGBA", "L"):
            image = image.convert("RGBA")
        out = BytesIO()
        image.save(out, "PNG")
        return out.getvalue(), image.width, image.height
    except Exception:
        return None


def _outline(lesson):
    """Hujjat tuzilmasi: sahnalar bo'yicha guruhlangan qadamlar."""
    groups = []
    for step in lesson.get("steps", []):
        if groups and groups[-1]["sahna"] == step.get("sahna"):
            groups[-1]["steps"].append(step)
        else:
            groups.append({"sahna": step.get("sahna"), "steps": [step]})
    return groups


def _meta(lesson):
    topic = lesson.get("topic") or {}
    rows = [("Fan", topic.get("fan")), ("Sinf", topic.get("sinf")),
            ("Daraja", f"{topic['daraja']} / 30" if topic.get("daraja") else ""),
            ("Kitob", (lesson.get("manba") or {}).get("kitob")), ("Dars maqsadi", topic.get("maqsad")),
            ("Mavzu kodi", topic.get("topic_code"))]
    return [(k, str(v)) for k, v in rows if v]


def lesson_docx(lesson, images):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    topic = lesson.get("topic") or {}
    head = doc.add_paragraph()
    run = head.add_run("OCHIQ DARS ISHLANMASI")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(10), RGBColor(0x2D, 0x8B, 0x8B)
    doc.add_heading(topic.get("mavzu") or topic.get("topic_code") or "Dars", level=0)
    meta = _meta(lesson)
    if meta:
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for key, value in meta:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = key, value
            cells[0].paragraphs[0].runs[0].bold = True
    doc.add_heading("Darsning borishi", level=1)
    number = 0
    for group in _outline(lesson):
        first = group["steps"][0]
        number += 1
        title = STEP_LABELS.get(first.get("turi"), "Qadam")
        if first.get("sarlavha"):
            title += f": {plain(first['sarlavha'])}"
        doc.add_heading(f"{number}. {title}", level=2)
        board = [plain(s.get("doska")) for s in group["steps"] if s.get("doska")]
        if board:
            p = doc.add_paragraph()
            p.add_run("Doskaga yoziladi: ").bold = True
            p.add_run("  →  ".join(board))
        for step in group["steps"]:
            if step.get("ovoz"):
                p = doc.add_paragraph()
                p.add_run("O'qituvchi: ").bold = True
                p.add_run(_voice(step["ovoz"]))
            picture = images.get(step.get("rasm") or "")
            converted = _png(picture) if picture else None
            if converted:
                doc.add_picture(BytesIO(converted[0]), width=Cm(min(9, converted[1] / 40)))
            if step.get("turi") == "birga" and step.get("javob"):
                p = doc.add_paragraph()
                p.add_run("O'quvchilar mustaqil bajaradi. To'g'ri javob: ").italic = True
                p.add_run(step["javob"].replace("|", " yoki "))
                if step.get("javob_izohi"):
                    doc.add_paragraph(step["javob_izohi"])
            variants = lesson.get("variants", {}).get(step.get("id")) or []
            if variants:
                p = doc.add_paragraph()
                p.add_run("Tushunmagan o'quvchilar uchun:").bold = True
                for v in variants:
                    item = doc.add_paragraph(style="List Bullet")
                    item.add_run(f"{VARIANT_LABELS.get(v.get('turi'), v.get('nom') or 'Variant')}: ").bold = True
                    item.add_run(" ".join(x for x in (plain(v.get("doska")), v.get("ovoz")) if x))
                    vpic = images.get(v.get("rasm") or "")
                    vconv = _png(vpic) if vpic else None
                    if vconv:
                        doc.add_picture(BytesIO(vconv[0]), width=Cm(min(7, vconv[1] / 40)))
    questions = lesson.get("savollar") or []
    if questions:
        doc.add_heading("Mustahkamlash testi", level=1)
        for i, q in enumerate(questions, 1):
            doc.add_paragraph(f"{i}. {plain(q['savol'])}")
            for k, option in enumerate(q["variantlar"]):
                doc.add_paragraph(f"{'ABCD'[k]}) {plain(option)}").paragraph_format.left_indent = Cm(0.8)
        doc.add_heading("Javoblar kaliti", level=2)
        for i, q in enumerate(questions, 1):
            line = f"{i} — {'ABCD'[q['togri']]}"
            if q.get("izoh"):
                line += f". {plain(q['izoh'])}"
            doc.add_paragraph(line)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def lesson_pdf(lesson, images):
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    with _LOCK:
        for name, file in (("Kabutar", "DejaVuSans.ttf"), ("KabutarBold", "DejaVuSans-Bold.ttf")):
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(FONT_ROOT / file)))
        pdfmetrics.registerFontFamily("Kabutar", normal="Kabutar", bold="KabutarBold", italic="Kabutar", boldItalic="KabutarBold")
    ink, teal, grey = colors.HexColor("#17314b"), colors.HexColor("#2D8B8B"), colors.HexColor("#526578")
    st = {
        "eyebrow": ParagraphStyle("e", fontName="KabutarBold", fontSize=9, textColor=teal, leading=12, spaceAfter=4),
        "title": ParagraphStyle("t", fontName="KabutarBold", fontSize=20, leading=25, textColor=ink, spaceAfter=10),
        "h1": ParagraphStyle("h1", fontName="KabutarBold", fontSize=14, leading=18, textColor=ink, spaceBefore=12, spaceAfter=6),
        "h2": ParagraphStyle("h2", fontName="KabutarBold", fontSize=11.5, leading=15, textColor=teal, spaceBefore=9, spaceAfter=4, keepWithNext=True),
        "body": ParagraphStyle("b", fontName="Kabutar", fontSize=10.5, leading=15, textColor=ink, spaceAfter=4),
        "bullet": ParagraphStyle("u", fontName="Kabutar", fontSize=10, leading=14, textColor=ink, leftIndent=14, bulletIndent=4, spaceAfter=3),
        "meta": ParagraphStyle("m", fontName="Kabutar", fontSize=9.5, leading=13, textColor=ink),
    }

    def para(label, text, style="body"):
        body = escape(_no_emoji(text)).replace("\n", "<br/>")
        return Paragraph(f"<b>{escape(label)}</b>{body}" if label else body, st[style])

    def picture(data, max_w=240):
        converted = _png(data) if data else None
        if not converted:
            return None
        png, w, h = converted
        scale = min(1.0, max_w / w, 170 / h)
        return Image(BytesIO(png), width=w * scale, height=h * scale, hAlign="LEFT")

    topic = lesson.get("topic") or {}
    story = [Paragraph("OCHIQ DARS ISHLANMASI", st["eyebrow"]),
             Paragraph(escape(topic.get("mavzu") or topic.get("topic_code") or "Dars"), st["title"])]
    meta = _meta(lesson)
    if meta:
        table = Table([[Paragraph(f"<b>{escape(k)}</b>", st["meta"]), Paragraph(escape(v), st["meta"])] for k, v in meta],
                      colWidths=[110, 390])
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5DBE1")),
                                   ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF3F6")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(table)
    story.append(Paragraph("Darsning borishi", st["h1"]))
    for number, group in enumerate(_outline(lesson), 1):
        first = group["steps"][0]
        title = STEP_LABELS.get(first.get("turi"), "Qadam")
        if first.get("sarlavha"):
            title += f": {plain(first['sarlavha'])}"
        story.append(Paragraph(escape(f"{number}. {title}"), st["h2"]))
        board = [plain(s.get("doska")) for s in group["steps"] if s.get("doska")]
        if board:
            story.append(para("Doskaga yoziladi: ", "  →  ".join(board)))
        for step in group["steps"]:
            if step.get("ovoz"):
                story.append(para("O'qituvchi: ", _voice(step["ovoz"])))
            img = picture(images.get(step.get("rasm") or ""))
            if img:
                story.append(img)
            if step.get("turi") == "birga" and step.get("javob"):
                story.append(para("To'g'ri javob: ", step["javob"].replace("|", " yoki ") + (f". {step['javob_izohi']}" if step.get("javob_izohi") else "")))
            variants = lesson.get("variants", {}).get(step.get("id")) or []
            if variants:
                story.append(para("Tushunmagan o'quvchilar uchun:", ""))
                for v in variants:
                    label = VARIANT_LABELS.get(v.get("turi"), v.get("nom") or "Variant")
                    text = " ".join(x for x in (plain(v.get("doska")), v.get("ovoz")) if x)
                    story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(_no_emoji(text))}", st["bullet"], bulletText="•"))
                    vimg = picture(images.get(v.get("rasm") or ""), 180)
                    if vimg:
                        story.append(vimg)
    questions = lesson.get("savollar") or []
    if questions:
        story.append(Paragraph("Mustahkamlash testi", st["h1"]))
        for i, q in enumerate(questions, 1):
            story.append(Paragraph(escape(f"{i}. {plain(q['savol'])}"), st["body"]))
            story.append(Paragraph("&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;".join(escape(f"{'ABCD'[k]}) {plain(o)}") for k, o in enumerate(q["variantlar"])), st["bullet"]))
        story.append(Spacer(1, 6))
        story.append(Paragraph("Javoblar kaliti", st["h2"]))
        for i, q in enumerate(questions, 1):
            story.append(Paragraph(escape(f"{i} — {'ABCD'[q['togri']]}" + (f". {plain(q['izoh'])}" if q.get("izoh") else "")), st["body"]))

    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=46, rightMargin=46, topMargin=44, bottomMargin=44,
                            title=f"Ochiq dars: {topic.get('mavzu') or ''}", author="Kabutar")

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Kabutar", 8)
        canvas.setFillColor(grey)
        canvas.drawRightString(A4[0] - 46, 26, f"{document.page}")
        canvas.drawString(46, 26, "Kabutar · kitob asosidagi dars")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()

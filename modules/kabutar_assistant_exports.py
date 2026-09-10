"""Build separate question papers and answer keys from a saved test attempt.

The caller must load an authorized attempt from the database and check answer-key
permissions before calling this module. No network or database access happens here.
"""

from __future__ import annotations

from io import BytesIO
import math
import re
import warnings
from typing import Any


_LANG_TAG = re.compile(r"\[\s*/?\s*(?:uz|ru|en)\s*\]", re.IGNORECASE)
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_MEDIA = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_LABELS = {
    "easy": "Oson", "medium": "O‘rtacha", "hard": "Qiyin", "mixed": "Aralash",
    "test": "Mashq testi", "exam": "Imtihon", "monitoring": "Monitoring",
    "single_choice": "Bitta javob", "write_answer": "Yozma javob",
}
_MAX_IMAGE_BYTES = 4 * 1024 * 1024
_MAX_TOTAL_IMAGE_BYTES = 20 * 1024 * 1024
_MAX_IMAGES = 100
_MAX_IMAGE_PIXELS = 16_000_000


def _text(value: Any) -> str:
    if value is None:
        return ""
    return _INVALID_XML.sub("", _LANG_TAG.sub("", str(value))).strip()


def _label(value: Any) -> str:
    value = _text(value)
    return _LABELS.get(value, value)


def _points(question: dict) -> float | int:
    try:
        value = float(question.get("points", 1))
    except (TypeError, ValueError):
        return 1
    if not math.isfinite(value) or value < 0:
        return 1
    return int(value) if value.is_integer() else value


def _topics(plan: dict) -> dict:
    topics = plan.get("topics") or []
    return {
        _text(topic.get("topic_code")): topic
        for topic in topics
        if isinstance(topic, dict)
    }


def _topic(question: dict, topics: dict) -> str:
    code = _text(question.get("topic_code"))
    item = topics.get(code) or {}
    return " · ".join(filter(None, (_text(item.get("subject_name")), _text(item.get("title"))))) or code


def _answer(question: dict) -> str:
    raw = _text(question.get("correct_answer"))
    if not raw:
        return "Javob kaliti bazada ko‘rsatilmagan"
    if question.get("question_type") == "write_answer":
        return raw
    if raw.upper() in ("A", "B", "C", "D"):
        return raw.upper()
    # Imported banks may store the option text instead of its letter.
    matches = [letter for letter in "ABCD" if _text(question.get("option_" + letter.lower())).casefold() == raw.casefold()]
    return f"{matches[0]}) {raw}" if len(matches) == 1 else raw


def _prepare_images(questions: list[dict]) -> dict[int, tuple[bytes, int, int]]:
    """Validate DB-provided images; never produce an incomplete question paper.

    Only a byte payload is accepted. URLs, file paths and SVG are deliberately
    unsupported. Images are decoded and re-encoded to PNG with bounded dimensions
    to remove metadata and keep document generation within a small memory budget.
    """
    from PIL import Image, UnidentifiedImageError

    pending = []
    missing = []
    total = 0
    for index, question in enumerate(questions, 1):
        raw = question.get("image_bytes")
        ident = _text(question.get("id"))[:80] or str(index)
        raw_size = raw.nbytes if isinstance(raw, memoryview) else len(raw) if isinstance(raw, (bytes, bytearray)) else None
        if raw is None or raw_size == 0:
            if question.get("rasm_id"):
                missing.append(ident)
            continue
        if not isinstance(raw, (bytes, bytearray, memoryview)):
            raise ValueError(f"Savol ID {ident}: rasm bazadan bayt ko‘rinishida olinishi kerak.")
        size = raw_size
        if size > _MAX_IMAGE_BYTES:
            raise ValueError(f"Savol ID {ident}: rasm 4 MB dan katta.")
        total += size
        if total > _MAX_TOTAL_IMAGE_BYTES:
            raise ValueError("Test rasmlarining jami hajmi 20 MB dan oshdi; testni qismlarga ajrating.")
        pending.append((index, ident, raw))
    if missing:
        raise ValueError("Rasm olinmadi. Savol ID: " + ", ".join(missing) + ". Eksport uchun rasmlarni tiklang.")
    if len(pending) > _MAX_IMAGES:
        raise ValueError("Bitta eksportda ko‘pi bilan 100 ta rasm bo‘lishi mumkin.")

    images = {}
    normalized_total = 0
    for index, ident, raw in pending:
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(BytesIO(raw)) as source:
                    if source.format not in {"PNG", "JPEG", "WEBP", "GIF"}:
                        raise ValueError(f"Savol ID {ident}: PNG, JPEG, WEBP yoki GIF rasm kerak.")
                    if source.width * source.height > _MAX_IMAGE_PIXELS or max(source.size) > 10000:
                        raise ValueError(f"Savol ID {ident}: rasm o‘lchamlari juda katta.")
                    source.verify()
                with Image.open(BytesIO(raw)) as source:
                    source.seek(0)
                    source.load()
                    normalized = source.convert("RGB" if source.format == "JPEG" else "RGBA")
                    normalized.thumbnail((1600, 1200), Image.Resampling.LANCZOS)
                    stream = BytesIO()
                    normalized.save(stream, format="PNG")
                    payload = stream.getvalue()
                    width, height = normalized.size
                    normalized.close()
            normalized_total += len(payload)
            if len(payload) > _MAX_IMAGE_BYTES or normalized_total > _MAX_TOTAL_IMAGE_BYTES:
                raise ValueError("Eksport rasmlari hajmi limitdan oshdi; testni qismlarga ajrating.")
            images[index] = (payload, width, height)
        except (UnidentifiedImageError, OSError, Image.DecompressionBombWarning, Image.DecompressionBombError) as exc:
            raise ValueError(f"Savol ID {ident}: rasm buzilgan yoki qo‘llab-quvvatlanmaydi.") from exc
    return images


def _fit_image(width: int, height: int, max_width: float, max_height: float) -> tuple[float, float]:
    scale = min(max_width / width, max_height / height)
    return width * scale, height * scale


def _metadata(attempt: dict, questions: list[dict]) -> list[tuple[str, str]]:
    plan = attempt.get("plan") or {}
    return [
        ("Test ID", _text(attempt.get("attempt_id"))),
        ("Sinf", _text(plan.get("grade"))),
        ("Savollar", str(len(questions))),
        ("Qiyinlik", _label(plan.get("difficulty"))),
        ("Vaqt (daqiqa)", _text(plan.get("minutes"))),
        ("Tartib", _label(plan.get("mode"))),
        ("Jami ball", str(sum(_points(question) for question in questions))),
    ]


def _docx(attempt: dict, questions: list[dict], answer_key: bool, images: dict) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1
    doc.styles["Title"].font.color.rgb = RGBColor.from_string("153E5C")
    doc.core_properties.title = "Kabutar — javoblar kaliti" if answer_key else "Kabutar — test savollari"
    doc.core_properties.author = "Kabutar"
    section.header.paragraphs[0].text = "KABUTAR · TA’LIM"
    doc.add_heading("Javoblar kaliti" if answer_key else "Test savollari", 0)
    for label, value in _metadata(attempt, questions):
        if value:
            paragraph = doc.add_paragraph()
            paragraph.add_run(label + ": ").bold = True
            paragraph.add_run(value)

    if answer_key:
        doc.add_paragraph("Ushbu fayl faqat javoblar kaliti uchun. O‘quvchiga beriladigan savollar alohida faylda.")
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, ("№", "Savol ID", "To‘g‘ri javob", "Ball")):
            cell.text = value
        for index, question in enumerate(questions, 1):
            for cell, value in zip(table.add_row().cells, (str(index), _text(question.get("id")), _answer(question), str(_points(question)))):
                cell.text = value
        explanations = [(index, _text(question.get("explanation"))) for index, question in enumerate(questions, 1) if _text(question.get("explanation"))]
        if explanations:
            doc.add_heading("Izohlar", 1)
            for index, explanation in explanations:
                doc.add_paragraph(f"{index}. {explanation}")
    else:
        doc.add_paragraph("Ism-familiya: __________________________________   Sana: ______________")
        topics = _topics(attempt.get("plan") or {})
        for index, question in enumerate(questions, 1):
            heading = doc.add_paragraph()
            heading.paragraph_format.keep_with_next = True
            heading.add_run(f"{index}-savol · {_points(question)} ball").bold = True
            title = _topic(question, topics)
            if title:
                doc.add_paragraph(title, style="Caption")
            doc.add_paragraph(_text(question.get("question")))
            if index in images:
                image, width, height = images[index]
                display_width, display_height = _fit_image(width, height, 15.6, 9.0)
                doc.add_picture(BytesIO(image), width=Cm(display_width), height=Cm(display_height))
            if question.get("question_type") == "write_answer":
                doc.add_paragraph("Javob: ______________________________________________________")
            else:
                for letter in "ABCD":
                    value = _text(question.get("option_" + letter.lower()))
                    if value:
                        doc.add_paragraph(f"{letter}) {value}")

    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _xlsx(attempt: dict, questions: list[dict], answer_key: bool, images: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SheetImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Javoblar kaliti" if answer_key else "Savollar"
    book.properties.creator = "Kabutar"
    book.properties.title = sheet.title
    fill = PatternFill("solid", fgColor="153E5C")

    def text_cell(target, row: int, col: int, value: Any):
        value = _text(value)
        if len(value) > 32767:
            raise ValueError("Excel katagiga sig‘maydigan matn bor; Word formatini tanlang.")
        cell = target.cell(row=row, column=col, value=value)
        # Binding a leading '=' normally turns it into an Excel formula.
        # Explicit string type keeps ALL database/user strings inert, including
        # names, IDs, answers, options and metadata, without changing their text.
        cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        return cell

    def header(target, labels):
        for col, label in enumerate(labels, 1):
            cell = text_cell(target, 1, col, label)
            cell.fill = fill
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        target.freeze_panes = "A2"
        target.sheet_view.showGridLines = False

    if answer_key:
        header(sheet, ("№", "Savol ID", "To‘g‘ri javob", "Izoh", "Ball"))
        for row, question in enumerate(questions, 2):
            sheet.cell(row, 1, row - 1)
            text_cell(sheet, row, 2, question.get("id"))
            text_cell(sheet, row, 3, _answer(question))
            text_cell(sheet, row, 4, question.get("explanation"))
            sheet.cell(row, 5, _points(question))
        widths = [7, 20, 38, 85, 9]
    else:
        header(sheet, ("№", "Savol ID", "Fan va mavzu", "Savol", "A", "B", "C", "D", "Savol turi", "Ball", "Rasm"))
        topics = _topics(attempt.get("plan") or {})
        for row, question in enumerate(questions, 2):
            sheet.cell(row, 1, row - 1)
            for col, value in enumerate((question.get("id"), _topic(question, topics), question.get("question")), 2):
                text_cell(sheet, row, col, value)
            for col, letter in enumerate("ABCD", 5):
                text_cell(sheet, row, col, question.get("option_" + letter.lower()) if question.get("question_type") != "write_answer" else "")
            text_cell(sheet, row, 9, _label(question.get("question_type")))
            sheet.cell(row, 10, _points(question))
            text_cell(sheet, row, 11, "")
            if row - 1 in images:
                payload, width, height = images[row - 1]
                image = SheetImage(BytesIO(payload))
                image.width, image.height = _fit_image(width, height, 350, 240)
                sheet.add_image(image, f"K{row}")
                sheet.row_dimensions[row].height = (image.height + 16) * 0.75
        widths = [7, 18, 32, 70, 30, 30, 30, 30, 18, 9, 55]

    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    info = book.create_sheet("Test haqida")
    header(info, ("Ma’lumot", "Qiymat"))
    for row, (label, value) in enumerate(_metadata(attempt, questions), 2):
        text_cell(info, row, 1, label)
        text_cell(info, row, 2, value)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 70
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def export_attempt(attempt: dict, format: str, answer_key: bool = False) -> tuple[bytes, str, str]:
    """Return file bytes, MIME type and safe filename for one authorized attempt.

    ``answer_key=False`` never writes correct answers or explanations. A true
    value produces a distinct key file; it does not append keys to a paper.
    Questions and their original order must come from the saved attempt so that
    separately downloaded keys continue to match the question numbering.
    For papers, each question with ``rasm_id`` must also have ``image_bytes``
    loaded directly from the database. Answer keys do not require images.
    """
    if format not in _MEDIA:
        raise ValueError("Format docx yoki xlsx bo‘lishi kerak.")
    if not isinstance(attempt, dict) or not isinstance(attempt.get("plan", {}), dict):
        raise ValueError("Test ma’lumotlari noto‘g‘ri.")
    questions = attempt.get("questions")
    if not isinstance(questions, list) or not questions or not all(isinstance(item, dict) for item in questions):
        raise ValueError("Eksport uchun saqlangan savollar topilmadi.")
    ident = re.sub(r"[^a-zA-Z0-9_-]", "", _text(attempt.get("attempt_id")))[:64] or "test"
    filename = f"kabutar_{ident}_{'javoblar' if answer_key else 'savollar'}.{format}"
    images = {} if answer_key else _prepare_images(questions)
    data = _docx(attempt, questions, answer_key, images) if format == "docx" else _xlsx(attempt, questions, answer_key, images)
    return data, _MEDIA[format], filename
